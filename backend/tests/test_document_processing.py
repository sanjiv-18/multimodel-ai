import asyncio
import json
from unittest.mock import AsyncMock, patch

import pytest
from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.agents.tutor_agent import GroundedTutorAgent, is_valid_tutor_response
from app.core.llm_provider import LLMProvider
from app.agents.ingestion_agent import MultimodalIngestionAgent
from app.ingestion.pdf_processor import (
    DocumentProcessingError,
    PDFProcessor,
    detect_file_type,
)
from app.models.db_models import Base, Course, DocumentChunk, Material
from app.rag.embeddings import EmbeddingService
from app.rag.retriever import CourseRetriever
from app.rag.vector_store import VectorStore


def make_pdf(text: str = "") -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode())
        output.extend(body + b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010} 00000 n \n".encode())
    output.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n".encode()
    )
    return bytes(output)


@pytest.fixture
def db_session():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    yield session
    session.close()
    Base.metadata.drop_all(engine)
    engine.dispose()


def test_text_pdf_extracts_readable_page_text_without_pdf_objects(tmp_path):
    source = "Binary search requires a sorted list and repeatedly halves the search interval."
    path = tmp_path / "lesson.pdf"
    path.write_bytes(make_pdf(source))

    chunks = PDFProcessor.extract_chunks(str(path), "course", "material", "lesson.pdf")

    assert chunks
    assert all(chunk["page_number"] == 1 for chunk in chunks)
    assert all(chunk["material_type"] == "pdf" for chunk in chunks)
    extracted = "\n".join(chunk["content"] for chunk in chunks)
    assert "Binary search requires a sorted list" in extracted
    assert "%PDF-" not in extracted
    assert "obj" not in extracted


def test_docx_extracts_paragraphs_and_tables(tmp_path):
    document = Document()
    document.add_paragraph(
        "Photosynthesis converts light energy into chemical energy in green plants."
    )
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Chlorophyll"
    table.cell(0, 1).text = "Absorbs light energy"
    path = tmp_path / "lesson.docx"
    document.save(path)

    chunks = PDFProcessor.extract_chunks(str(path), "course", "material", "lesson.docx")

    content = "\n".join(chunk["content"] for chunk in chunks)
    assert "Photosynthesis converts light energy" in content
    assert "Chlorophyll | Absorbs light energy" in content
    assert all(chunk["material_type"] == "docx" for chunk in chunks)
    assert all(chunk["page_number"] is None for chunk in chunks)


def test_text_file_is_validated_and_has_no_fabricated_page_reference(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text(
        "Cellular respiration releases energy from glucose through a series of reactions.",
        encoding="utf-8",
    )

    chunks = PDFProcessor.extract_chunks(str(path), "course", "material", "notes.txt")

    assert chunks[0]["page_number"] is None
    assert "Cellular respiration" in chunks[0]["content"]


def test_invalid_pdf_is_not_decoded_as_fallback_text(tmp_path):
    path = tmp_path / "corrupt.pdf"
    path.write_bytes(b"%PDF-1.4\n1 0 obj\nstream\n\x00\xffcompressed bytes\nendstream")

    with pytest.raises(DocumentProcessingError, match="invalid, corrupted"):
        PDFProcessor.extract_chunks(str(path), "course", "material", "corrupt.pdf")


def test_ingestion_indexes_pdf_text_and_marks_corrupt_documents_failed(db_session, tmp_path):
    course = Course(title="Algorithms")
    db_session.add(course)
    db_session.flush()
    valid_path = tmp_path / "Algorithms.pdf"
    valid_path.write_bytes(make_pdf(
        "Binary search requires a sorted list and repeatedly halves the search interval."
    ))
    material = Material(
        course_id=course.id,
        title="Algorithms.pdf",
        file_type="pdf",
        file_path=str(valid_path),
        file_size_bytes=valid_path.stat().st_size,
        status="uploaded",
    )
    db_session.add(material)
    db_session.commit()

    result = MultimodalIngestionAgent(db_session).process_material(material.id)

    db_session.refresh(material)
    stored = db_session.query(DocumentChunk).filter_by(material_id=material.id).all()
    assert result["status"] == "completed"
    assert material.status == "completed"
    assert stored and "Binary search requires a sorted list" in stored[0].content
    assert all("%PDF-" not in chunk.content for chunk in stored)
    assert stored[0].source_name == "Algorithms.pdf"
    assert stored[0].page_number == 1

    corrupt_path = tmp_path / "broken.pdf"
    corrupt_path.write_bytes(b"%PDF-1.4\nnot a readable document")
    broken = Material(
        course_id=course.id,
        title="broken.pdf",
        file_type="pdf",
        file_path=str(corrupt_path),
        file_size_bytes=corrupt_path.stat().st_size,
        status="uploaded",
    )
    db_session.add(broken)
    db_session.commit()

    failed_result = MultimodalIngestionAgent(db_session).process_material(broken.id)

    db_session.refresh(broken)
    assert failed_result["status"] == "error"
    assert broken.status == "error"
    assert broken.error_message
    assert db_session.query(DocumentChunk).filter_by(material_id=broken.id).count() == 0


def test_scanned_pdf_reports_ocr_limitation(tmp_path):
    path = tmp_path / "scan.pdf"
    path.write_bytes(make_pdf())

    with pytest.raises(DocumentProcessingError, match="scanned.*OCR is not configured"):
        PDFProcessor.extract_chunks(str(path), "course", "material", "scan.pdf")


def test_upload_format_detection_rejects_misnamed_and_unsupported_files():
    with pytest.raises(DocumentProcessingError, match="does not match"):
        detect_file_type("lesson.pdf", "application/pdf", b"plain text renamed as a PDF")
    with pytest.raises(DocumentProcessingError, match="Unsupported file type"):
        detect_file_type("lesson.exe", "application/octet-stream", b"MZ\x00\x01")
    assert detect_file_type("lesson.pdf", "application/pdf", make_pdf("Readable study text.")) == "pdf"


def _add_material_chunk(db_session, course, content, source, page):
    material = Material(
        course_id=course.id,
        title=source,
        file_type="pdf",
        file_path=source,
        file_size_bytes=1,
        status="completed",
    )
    db_session.add(material)
    db_session.flush()
    chunk = {
        "course_id": course.id,
        "material_id": material.id,
        "material_type": "pdf",
        "source_name": source,
        "page_number": page,
        "topic": "Binary Search",
        "concept": "Prerequisite",
        "content": content,
    }
    VectorStore(db_session).add_chunks([chunk])
    return material


def test_retrieval_is_scoped_and_citation_keeps_source_page(db_session):
    course = Course(title="Algorithms")
    db_session.add(course)
    db_session.flush()
    first = _add_material_chunk(
        db_session,
        course,
        "Binary search only works on a sorted list. Each comparison removes half the remaining values.",
        "Algorithms.pdf",
        7,
    )
    other = _add_material_chunk(
        db_session,
        course,
        "Binary search requires a sorted list and divides the search interval in half.",
        "Other.pdf",
        12,
    )

    context, citations, supported = CourseRetriever(db_session).retrieve_grounded_context(
        course.id,
        "Why does binary search require sorted values?",
        material_id=first.id,
    )

    assert supported
    assert len(citations) == 1
    assert citations[0].source_name == "Algorithms.pdf"
    assert citations[0].page_number == 7
    assert "Other.pdf" not in "\n".join(context)
    assert db_session.query(DocumentChunk).filter_by(material_id=other.id).count() == 1


def test_legacy_pdf_binary_chunks_are_never_retrieved(db_session):
    course = Course(title="Algorithms")
    db_session.add(course)
    db_session.flush()
    material = Material(
        course_id=course.id,
        title="OldUpload.pdf",
        file_type="pdf",
        file_path="OldUpload.pdf",
        file_size_bytes=100,
        status="completed",
    )
    db_session.add(material)
    db_session.flush()
    raw_pdf = "%PDF-1.4 1 0 obj stream compressed bytes endstream endobj"
    db_session.add(DocumentChunk(
        course_id=course.id,
        material_id=material.id,
        content=raw_pdf,
        source_name="OldUpload.pdf",
        material_type="pdf",
        page_number=1,
        embedding_json=json.dumps(EmbeddingService.get_embedding_vector(raw_pdf)),
    ))
    db_session.commit()

    context, citations, supported = CourseRetriever(db_session).retrieve_grounded_context(
        course.id, "What is in this binary PDF?"
    )

    assert not supported
    assert context == []
    assert citations == []


def test_grounded_and_simplified_answers_use_retrieved_material(db_session):
    course = Course(title="Algorithms")
    db_session.add(course)
    db_session.flush()
    material = _add_material_chunk(
        db_session,
        course,
        "Binary search requires a sorted list. It checks the middle value and removes half of the remaining list.",
        "Algorithms.pdf",
        7,
    )
    agent = GroundedTutorAgent(db_session)
    generated = AsyncMock(side_effect=[
        "Binary search needs sorted values because it discards half of the list after checking the middle.",
        "Think of a sorted list like a dictionary: opening near the middle lets you skip half the pages each time.",
    ])

    with patch("app.agents.tutor_agent.LLMProvider.generate_response", generated):
        first = asyncio.run(agent.answer_query(
            course.id, "student", "Why does binary search require sorted values?",
            material_id=material.id,
        ))
        follow_up = asyncio.run(agent.answer_query(
            course.id, "student",
            "Can you explain this again in simpler terms with a basic analogy?",
            conversation_id=first["conversation_id"],
            material_id=material.id,
        ))

    assert first["message"]["is_grounded"]
    assert first["message"]["citations"][0]["source_name"] == "Algorithms.pdf"
    assert first["message"]["citations"][0]["page_number"] == 7
    assert "dictionary" in follow_up["message"]["content"]
    follow_up_prompt = generated.await_args_list[1].args[0]
    assert "Binary search requires a sorted list" in follow_up_prompt
    assert "Can you explain this again in simpler terms" in follow_up_prompt


def test_unsupported_question_refuses_without_generating_or_citing(db_session):
    course = Course(title="Algorithms")
    db_session.add(course)
    db_session.flush()
    _add_material_chunk(
        db_session,
        course,
        "Binary search requires a sorted list.",
        "Algorithms.pdf",
        7,
    )
    generated = AsyncMock()

    with patch("app.agents.tutor_agent.LLMProvider.generate_response", generated):
        result = asyncio.run(GroundedTutorAgent(db_session).answer_query(
            course.id,
            "student",
            "How do I bake a chocolate cake with icing?",
        ))

    assert not result["message"]["is_grounded"]
    assert result["message"]["citations"] == []
    assert "couldn't find" in result["message"]["content"].lower()
    generated.assert_not_awaited()


def test_malformed_model_output_is_retried_and_never_returned(db_session):
    course = Course(title="Algorithms")
    db_session.add(course)
    db_session.flush()
    _add_material_chunk(
        db_session,
        course,
        "Binary search requires a sorted list and halves the remaining search interval.",
        "Algorithms.pdf",
        7,
    )
    generated = AsyncMock(side_effect=[
        "From [Algorithms.pdf — Page 7]: %PDF-1.4 1 0 obj stream \x00",
        "Binary search checks the middle of a sorted list and discards the half that cannot contain the target.",
    ])

    with patch("app.agents.tutor_agent.LLMProvider.generate_response", generated):
        result = asyncio.run(GroundedTutorAgent(db_session).answer_query(
            course.id,
            "student",
            "Explain binary search on a sorted list.",
        ))

    assert generated.await_count == 2
    assert "%PDF" not in result["message"]["content"]
    assert result["message"]["is_grounded"]
    assert is_valid_tutor_response(result["message"]["content"])


def test_response_validator_rejects_pdf_internals_and_empty_output():
    assert not is_valid_tutor_response("")
    assert not is_valid_tutor_response("%PDF-1.4 1 0 obj stream \x00")
    assert is_valid_tutor_response("Binary search checks the middle of a sorted list.")
    assert is_valid_tutor_response('A PDF syntax example may contain "1 0 obj".')


def test_builtin_tutor_fallback_reads_the_validated_context_header():
    answer = LLMProvider._fallback_handler(
        "COURSE MATERIAL EXCERPTS:\n"
        "[Algorithms.pdf — Page 7]\nBinary search requires a sorted list.\n\n"
        "STUDENT QUESTION:\nWhy must the list be sorted?",
        "Answer using the supplied course material.",
    )

    assert "Binary search requires a sorted list." in answer
    assert "Based on the analyzed course material" not in answer
