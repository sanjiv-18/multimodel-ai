import os
import re
import zipfile
from io import BytesIO
from typing import Any, Dict, List, Optional


class DocumentProcessingError(ValueError):
    """A safe, user-facing error raised when a document cannot be indexed."""


PDF_HEADER = re.compile(r"%PDF-\d")
PDF_OBJECT = re.compile(r"\b\d+\s+\d+\s+obj\b")


def is_readable_content(text: str) -> bool:
    normalized = text.strip()
    object_declarations = len(PDF_OBJECT.findall(normalized))
    object_terminators = len(re.findall(r"\bendobj\b|\bendstream\b", normalized))
    if (
        len(normalized) < 10
        or PDF_HEADER.search(normalized)
        or object_declarations >= 2
        or (object_declarations == 1 and object_terminators >= 2)
    ):
        return False
    visible = sum(character.isprintable() or character in "\n\r\t" for character in normalized)
    useful = sum(character.isalnum() for character in normalized)
    return visible / len(normalized) >= 0.85 and useful / len(normalized) >= 0.2


def _paragraphs(text: str) -> List[str]:
    return [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]


def _split_long_text(text: str, limit: int = 1200) -> List[str]:
    if len(text) <= limit:
        return [text]
    sentences = re.split(r"(?<=[.!?])\s+", text)
    pieces: List[str] = []
    current = ""
    for sentence in sentences:
        while len(sentence) > limit:
            if current:
                pieces.append(current)
                current = ""
            boundary = sentence.rfind(" ", 0, limit)
            boundary = boundary if boundary > limit // 2 else limit
            pieces.append(sentence[:boundary].strip())
            sentence = sentence[boundary:].strip()
        candidate = f"{current} {sentence}".strip()
        if current and len(candidate) > limit:
            pieces.append(current)
            current = sentence
        else:
            current = candidate
    if current:
        pieces.append(current)
    return pieces


def _chunk_text(
    text: str,
    course_id: str,
    material_id: str,
    source_name: str,
    material_type: str,
    page_number: Optional[int] = None,
) -> List[Dict[str, Any]]:
    paragraphs = _paragraphs(text)
    grouped: List[str] = []
    current = ""
    for paragraph in paragraphs:
        for part in _split_long_text(paragraph):
            candidate = f"{current}\n\n{part}".strip()
            if current and len(candidate) > 1200:
                grouped.append(current)
                current = part
            else:
                current = candidate
    if current:
        grouped.append(current)

    chunks = []
    for content in grouped:
        if not is_readable_content(content):
            continue
        first_line = content.splitlines()[0].strip()
        topic = first_line[:60] if len(first_line) < 60 else "Core Concepts"
        chunks.append({
            "course_id": course_id,
            "material_id": material_id,
            "material_type": material_type,
            "source_name": source_name,
            "page_number": page_number,
            "slide_number": None,
            "video_timestamp": None,
            "topic": topic,
            "concept": topic,
            "content": content,
        })
    return chunks


class PDFProcessor:
    @staticmethod
    def extract_chunks(
        file_path: str,
        course_id: str,
        material_id: str,
        source_name: str,
        file_type: Optional[str] = None,
    ) -> List[Dict[str, Any]]:
        if not os.path.isfile(file_path):
            raise DocumentProcessingError("The uploaded document is no longer available.")

        doc_type = (file_type or os.path.splitext(file_path)[1].lstrip(".")).lower()
        try:
            with open(file_path, "rb") as document:
                data = document.read()
        except OSError as exc:
            raise DocumentProcessingError("The uploaded document could not be read.") from exc

        if doc_type == "pdf":
            return PDFProcessor._extract_pdf(data, course_id, material_id, source_name)
        if doc_type == "docx":
            return PDFProcessor._extract_docx(data, course_id, material_id, source_name)
        if doc_type in {"txt", "text", "md", "markdown"}:
            return PDFProcessor._extract_text(data, course_id, material_id, source_name)
        raise DocumentProcessingError("This document format is not supported for text extraction.")

    @staticmethod
    def _extract_pdf(
        data: bytes, course_id: str, material_id: str, source_name: str
    ) -> List[Dict[str, Any]]:
        if b"%PDF-" not in data[:1024]:
            raise DocumentProcessingError("The file does not contain a valid PDF signature.")

        try:
            import pypdf
            reader = pypdf.PdfReader(BytesIO(data), strict=True)
            pages = [
                (page_number, (page.extract_text() or "").strip())
                for page_number, page in enumerate(reader.pages, start=1)
            ]
        except Exception as exc:
            raise DocumentProcessingError(
                "The PDF is invalid, corrupted, or could not be parsed."
            ) from exc

        chunks: List[Dict[str, Any]] = []
        for page_number, text in pages:
            if is_readable_content(text):
                chunks.extend(_chunk_text(
                    text, course_id, material_id, source_name, "pdf", page_number
                ))
        if not chunks:
            raise DocumentProcessingError(
                "No readable text could be extracted from this PDF. It may be scanned; "
                "OCR is not configured, so upload a text-based PDF or provide OCR text."
            )
        return chunks

    @staticmethod
    def _extract_docx(
        data: bytes, course_id: str, material_id: str, source_name: str
    ) -> List[Dict[str, Any]]:
        try:
            from docx import Document
            document = Document(BytesIO(data))
        except Exception as exc:
            raise DocumentProcessingError(
                "The DOCX file is invalid, corrupted, or could not be parsed."
            ) from exc

        parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n\n".join(parts)
        if not is_readable_content(text):
            raise DocumentProcessingError("No readable paragraphs or table text were found in this DOCX.")
        return _chunk_text(text, course_id, material_id, source_name, "docx")

    @staticmethod
    def _extract_text(
        data: bytes, course_id: str, material_id: str, source_name: str
    ) -> List[Dict[str, Any]]:
        try:
            if data.startswith((b"\xff\xfe", b"\xfe\xff")):
                text = data.decode("utf-16")
            else:
                text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DocumentProcessingError(
                "The text file is not valid UTF-8 or UTF-16 text."
            ) from exc
        if not is_readable_content(text):
            raise DocumentProcessingError("The text file is empty or does not contain readable study material.")
        return _chunk_text(text, course_id, material_id, source_name, "text")


def detect_file_type(filename: str, content_type: Optional[str], data: bytes) -> str:
    extension = os.path.splitext(filename.replace("\\", "/"))[1].lower()
    declared_mime = (content_type or "").split(";", 1)[0].strip().lower()
    detected: Optional[str] = None

    if b"%PDF-" in data[:1024]:
        detected = "pdf"
    elif data.startswith(b"PK"):
        try:
            with zipfile.ZipFile(BytesIO(data)) as archive:
                names = set(archive.namelist())
                if "word/document.xml" in names and "[Content_Types].xml" in names:
                    detected = "docx"
                elif "ppt/presentation.xml" in names and "[Content_Types].xml" in names:
                    detected = "pptx"
        except (OSError, zipfile.BadZipFile):
            detected = None
    elif data.startswith(b"RIFF") and data[8:12] == b"AVI ":
        detected = "video"
    elif data.startswith(b"\x1aE\xdf\xa3") or (len(data) >= 8 and data[4:8] == b"ftyp"):
        detected = "video"
    elif data.startswith(b"ID3") or (len(data) >= 2 and data[0] == 0xFF and data[1] & 0xE0 == 0xE0):
        detected = "video"
    elif data.startswith(b"RIFF") and data[8:12] == b"WAVE":
        detected = "video"

    text_extensions = {".txt", ".md", ".markdown"}
    extension_types = {
        ".pdf": "pdf",
        ".docx": "docx",
        ".pptx": "pptx",
        ".txt": "text",
        ".md": "text",
        ".markdown": "text",
        ".mp4": "video",
        ".mkv": "video",
        ".avi": "video",
        ".mov": "video",
        ".mp3": "video",
        ".wav": "video",
    }
    expected = extension_types.get(extension)
    if expected is None:
        raise DocumentProcessingError(
            "Unsupported file type. Supported formats are PDF, DOCX, PPTX, TXT, MD, and common audio/video files."
        )

    if expected == "text":
        try:
            text = data.decode("utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DocumentProcessingError("The text file is not valid UTF-8 or UTF-16 text.") from exc
        if not is_readable_content(text):
            raise DocumentProcessingError("The text file is empty or does not contain readable study material.")
        if detected is not None:
            raise DocumentProcessingError("The file content does not match its .txt or .md extension.")
    elif detected != expected:
        raise DocumentProcessingError(
            f"The file content does not match its .{extension.lstrip('.')} extension."
        )

    recognized_mimes = {
        "pdf": {"application/pdf"},
        "docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
        "pptx": {"application/vnd.openxmlformats-officedocument.presentationml.presentation"},
        "text": {"text/plain", "text/markdown"},
        "video": {
            "video/mp4", "video/x-matroska", "video/x-msvideo", "video/quicktime",
            "audio/mpeg", "audio/wav", "audio/x-wav",
        },
    }
    mime_expected = next(
        (file_type for file_type, mime_types in recognized_mimes.items() if declared_mime in mime_types),
        None,
    )
    if declared_mime and declared_mime != "application/octet-stream" and declared_mime not in recognized_mimes[expected]:
        raise DocumentProcessingError("The uploaded file content does not match its declared MIME type.")
    if mime_expected and mime_expected != expected:
        raise DocumentProcessingError("The uploaded file content does not match its declared MIME type.")
    return expected
